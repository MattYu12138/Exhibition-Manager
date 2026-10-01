#!/usr/bin/env python3
"""Reconcile LIC260001 trade documents; output non-financial private JSON to stdout.

This is *not* GOTS certification, a manufacturing batch, or a warehouse receipt.
Usage: python3 scripts/buildTradeManifest.py PO_WITH_BARCODES.xlsx INVOICE.xlsx PACK.xlsx BOL.pdf ORIGINAL_PO.xlsx > /private/manifest.json
"""
import sys
import json
import re
import hashlib
from pathlib import Path
import openpyxl
from pypdf import PdfReader


def sku(value):
    return re.sub(r'\s*-\s*', '-', str(value or '').strip()).upper()


def quantity(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0 or int(value) != value:
        raise ValueError(f'Invalid quantity in source document: {value!r}')
    return int(value)


def unique_lines(sheet, key_col, start_row, qty_col, title_col):
    result = {}
    for row in sheet.iter_rows(min_row=start_row):
        key = sku(row[key_col - 1].value)
        if not re.match(r'^[A-Z]{2,8}\d{4,8}(?:-[A-Z0-9]+)?$', key):
            continue
        if key in result:
            raise ValueError(f'Duplicate document SKU: {key}')
        qty = row[qty_col - 1].value
        result[key] = {'qty': quantity(qty), 'title': str(row[title_col - 1].value or '').strip()}
    return result


def build(paths):
    po_path, invoice_path, packing_path, bol_path, original_po_path = [Path(p) for p in paths]
    po, invoice, pack = [openpyxl.load_workbook(p, read_only=True, data_only=True).active
                         for p in (po_path, invoice_path, packing_path)]
    original = openpyxl.load_workbook(original_po_path, read_only=True, data_only=True).active
    if original.max_row != po.max_row or any(
        [cell.value for cell in raw_row] != [cell.value for cell in enriched_row]
        for raw_row, enriched_row in zip(original.iter_rows(max_col=9), po.iter_rows(max_col=9))
    ):
        raise ValueError('Barcode-enriched copy changed the original PO columns A-I')
    bill_text = ' '.join(page.extract_text() or '' for page in PdfReader(bol_path).pages).upper()
    for marker in ('SL0202302260M', '2026/07/30', 'NANSHA', 'MELBOURNE', '1039', '7.48', '102'):
        if marker not in bill_text:
            raise ValueError(f'Bill of lading missing expected shipment fact: {marker}')
    vessel = re.search(r'INTENDED VESSEL\s*&\s*VOY\s+TERM\s+([A-Z0-9 ]+?)\s+(DDU|DDP|FOB|CIF)\b', bill_text)
    container = re.search(r'\bNO\.?:\s*(XHCU\d{7})\b', bill_text)
    if not vessel or not container or vessel.group(1).strip() != 'MSC ODESSA V 29S':
        raise ValueError('Bill of lading vessel or container differs from the original document')
    for label, sheet, marker in (
        ('original PO', original, 'LIC260001'), ('invoice', invoice, 'UNA260001'),
        ('packing list', pack, 'UNA260001'),
    ):
        header = ' '.join(str(cell.value or '') for row in sheet.iter_rows(max_row=22) for cell in row)
        if marker not in header.upper():
            raise ValueError(f'{label} does not display reference {marker}')
    if (pack['H141'].value != 12060 or pack['I141'].value != 102
            or abs(float(pack['L141'].value or 0) - 957.2) > 0.01
            or abs(float(pack['M141'].value or 0) - 1038.8) > 0.01):
        raise ValueError('Packing-list totals differ from the supplied original')
    ordered = []
    po_by_sku = unique_lines(po, 1, 14, 6, 2)
    invoice_by_sku = unique_lines(invoice, 2, 23, 11, 9)
    packing_by_sku = {}
    for row in pack.iter_rows(min_row=15):
        key = sku(row[2].value)
        if not re.match(r'^[A-Z]{2,8}\d{4,8}(?:-[A-Z0-9]+)?$', key):
            continue
        if key in packing_by_sku:
            raise ValueError(f'Duplicate packing-list SKU: {key}')
        packing_by_sku[key] = {
            'units_per_carton_or_group': quantity(row[6].value),
            'carton_group_total': quantity(row[7].value) if row[7].value is not None else None,
        }
    if not (len(po_by_sku) == len(invoice_by_sku) == len(packing_by_sku) == 125
            and set(po_by_sku) == set(invoice_by_sku) == set(packing_by_sku)):
        raise ValueError('PO/invoice/packing SKU sets differ; review original documents before importing')
    for row in po.iter_rows(min_row=14):
        key = sku(row[0].value)
        if key not in po_by_sku:
            continue
        po_qty = po_by_sku[key]['qty']
        invoice_qty = invoice_by_sku[key]['qty']
        # Column G is per-carton for ordinary rows, but the exact SKU units for
        # multiple sizes sharing a carton. Column H is either the SKU's total or
        # the *whole grouped carton* total. Reconcile both against PO/invoice.
        candidates = (packing_by_sku[key]['carton_group_total'], packing_by_sku[key]['units_per_carton_or_group'])
        packing_qty = po_qty if po_qty in candidates else (candidates[0] or candidates[1])
        if po_qty != invoice_qty or po_qty != packing_qty:
            raise ValueError(f'Document quantity mismatch for {key}: PO {po_qty}; invoice {invoice_qty}; packing {packing_qty}')
        barcode = str(row[9].value or '').strip() if len(row) > 9 else ''
        if barcode and not re.fullmatch(r'\d{8}', barcode):
            raise ValueError(f'Invalid 8-digit barcode on purchase order for {key}')
        ordered.append({
            'sku': key, 'title': po_by_sku[key]['title'],
            'size': str(row[3].value or '').strip(), 'barcode': barcode,
            'po_quantity': po_qty, 'invoice_quantity': invoice_qty,
            'packing_quantity': packing_qty,
        })
    if sum(line['packing_quantity'] for line in ordered) != 12060:
        raise ValueError('Document units total differs from stated invoice/packing total of 12060')
    hashes = {kind: hashlib.sha256(p.read_bytes()).hexdigest() for kind, p in zip(
        ('original_purchase_order', 'barcode_enriched_working_copy', 'invoice', 'packing_list', 'bill_of_lading'),
        (original_po_path, po_path, invoice_path, packing_path, bol_path))}
    return {
        'id': 'TRADE-UNA260001', 'po_ref': 'LIC260001', 'invoice_ref': 'UNA260001',
        'packing_ref': 'UNA260001', 'bol_ref': 'SL0202302260M',
        'supplier_name': 'WUHAN U & MEE CO., LTD', 'shipped_at': '2026-07-30',
        'port_of_loading': 'Nansha, China', 'port_of_discharge': 'Melbourne, Australia',
        'declared_cartons': 102, 'declared_units': 12060,
        'intended_vessel_voyage': vessel.group(1).strip(), 'container_no': container.group(1),
        'delivery_term': vessel.group(2), 'bol_gross_weight_kg': 1039,
        'bol_measurement_cbm': 7.48, 'packing_net_weight_kg': float(pack['L141'].value),
        'packing_gross_weight_kg': float(pack['M141'].value),
        'source_checksums': hashes, 'lines': ordered,
        'disclaimer': 'Commercial and shipping documents only; no verified production lot, SC, TC, label approval or stock receipt link.',
    }


if __name__ == '__main__':
    if len(sys.argv) != 6:
        sys.exit('Usage: buildTradeManifest.py PO_WITH_BARCODES.xlsx INVOICE.xlsx PACK.xlsx BOL.pdf ORIGINAL_PO.xlsx')
    print(json.dumps(build(sys.argv[1:]), ensure_ascii=False, separators=(',', ':')))
